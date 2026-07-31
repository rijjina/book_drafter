from __future__ import annotations

import base64
import hashlib
import json
import subprocess
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

import yaml
from docx import Document


SKILL_ROOT = Path(__file__).resolve().parents[1]
BUILDER = SKILL_ROOT / "scripts" / "build_docx.py"
PREFLIGHT = SKILL_ROOT / "scripts" / "docx_preflight.py"

PNG_1X1 = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAusB9Y9ZQmcAAAAASUVORK5CYII="
)


def file_hash(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def run_json(*args: str) -> tuple[subprocess.CompletedProcess[str], dict]:
    completed = subprocess.run(
        [sys.executable, *args],
        capture_output=True,
        text=True,
        encoding="utf-8",
        check=False,
    )
    payload = json.loads(completed.stdout)
    return completed, payload


def write_manifest(root: Path, include_assets: bool = False) -> Path:
    manifest: dict = {
        "version": 1,
        "project": {
            "title": "หนังสือทดสอบ",
            "subject": "จุลชีววิทยาสิ่งแวดล้อม",
            "author": "ผู้เขียนทดสอบ",
            "keywords": ["จุลชีววิทยา", "สิ่งแวดล้อม"],
            "header_text": "หนังสือทดสอบ",
        },
        "build": {
            "directory": "build",
            "state": "build/state.json",
            "manuscript_output": "final/manuscript.docx",
        },
        "chapters": [
            {
                "id": "chapter-01",
                "source": "chapters/chapter-01/draft.md",
                "output": "chapters/chapter-01/chapter-01.docx",
                "metadata": {"title": "บทที่ 1 บทนำ"},
            },
            {
                "id": "chapter-02",
                "source": "chapters/chapter-02/draft.md",
                "output": "chapters/chapter-02/chapter-02.docx",
                "metadata": {"title": "บทที่ 2 เนื้อหา"},
            },
        ],
    }
    if include_assets:
        manifest["chapters"] = [manifest["chapters"][0]]
        manifest["chapters"][0]["assets"] = {
            "images": [
                {
                    "id": "framework",
                    "path": "assets/framework.png",
                    "alt": "กรอบแนวคิดทดสอบ",
                    "caption": "ภาพที่ 1.1 กรอบแนวคิดทดสอบ",
                    "width_cm": 2,
                }
            ],
            "tables": [
                {
                    "id": "comparison",
                    "path": "assets/comparison.csv",
                    "caption": "ตารางที่ 1.1 ตารางเปรียบเทียบทดสอบ",
                    "widths": [1, 2],
                }
            ],
        }
    path = root / "document.yml"
    path.write_text(yaml.safe_dump(manifest, allow_unicode=True, sort_keys=False), encoding="utf-8")
    return path


def write_manifest_v2(root: Path) -> Path:
    manifest = {
        "version": 2,
        "project": {"title": "หนังสือทดสอบ", "language": "th-TH"},
        "build": {
            "directory": ".docx-build",
            "state": ".docx-build/state.json",
            "manuscript_output": "final/manuscript.docx",
        },
        "chapters": [
            {
                "id": "chapter-01",
                "source_candidates": [
                    "chapters/chapter-01/revised.md",
                    "chapters/chapter-01/draft.md",
                ],
            }
        ],
    }
    path = root / "document.yml"
    path.write_text(
        yaml.safe_dump(manifest, allow_unicode=True, sort_keys=False),
        encoding="utf-8",
    )
    return path


class DeterministicBuilderTests(unittest.TestCase):
    def test_changed_only_rebuilds_only_the_changed_chapter(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            chapter_1 = root / "chapters" / "chapter-01" / "draft.md"
            chapter_2 = root / "chapters" / "chapter-02" / "draft.md"
            chapter_1.parent.mkdir(parents=True)
            chapter_2.parent.mkdir(parents=True)
            chapter_1.write_text(
                "# บทที่ 1 บทนำ\n\n## 1.1 หลักการ\n\nเนื้อหาบทแรกสำหรับการทดสอบ\n",
                encoding="utf-8",
            )
            chapter_2.write_text(
                "# บทที่ 2 เนื้อหา\n\n## 2.1 หลักการ\n\nเนื้อหาบทที่สองสำหรับการทดสอบ\n",
                encoding="utf-8",
            )
            manifest = write_manifest(root)

            first, first_payload = run_json(str(BUILDER), "--manifest", str(manifest))
            self.assertEqual(first.returncode, 0, first.stderr)
            self.assertEqual(first_payload["rebuilt_chapters"], ["chapter-01", "chapter-02"])
            output_1 = root / "chapters" / "chapter-01" / "chapter-01.docx"
            output_2 = root / "chapters" / "chapter-02" / "chapter-02.docx"
            first_hash_1 = file_hash(output_1)
            first_hash_2 = file_hash(output_2)

            repeat, repeat_payload = run_json(str(BUILDER), "--manifest", str(manifest))
            self.assertEqual(repeat.returncode, 0, repeat.stderr)
            self.assertEqual(first_hash_1, file_hash(output_1))
            self.assertEqual(first_hash_2, file_hash(output_2))
            self.assertEqual(repeat_payload["rebuilt_chapters"], ["chapter-01", "chapter-02"])

            skipped, skipped_payload = run_json(
                str(BUILDER), "--manifest", str(manifest), "--changed-only"
            )
            self.assertEqual(skipped.returncode, 0, skipped.stderr)
            self.assertEqual(skipped_payload["status"], "SKIPPED")
            self.assertEqual(skipped_payload["skipped_chapters"], ["chapter-01", "chapter-02"])

            chapter_1.write_text(
                "# บทที่ 1 บทนำ\n\n## 1.1 หลักการ\n\nเนื้อหาบทแรกที่แก้ไขเฉพาะบทนี้\n",
                encoding="utf-8",
            )
            changed, changed_payload = run_json(
                str(BUILDER), "--manifest", str(manifest), "--changed-only"
            )
            self.assertEqual(changed.returncode, 0, changed.stderr)
            self.assertEqual(changed_payload["rebuilt_chapters"], ["chapter-01"])
            self.assertEqual(changed_payload["skipped_chapters"], ["chapter-02"])
            self.assertNotEqual(first_hash_1, file_hash(output_1))
            self.assertEqual(first_hash_2, file_hash(output_2))

    def test_assembly_is_explicit_and_incremental(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            for number in (1, 2):
                source = root / "chapters" / f"chapter-{number:02d}" / "draft.md"
                source.parent.mkdir(parents=True)
                source.write_text(
                    f"# บทที่ {number}\n\n## {number}.1 หัวข้อ\n\nเนื้อหาทดสอบบทที่ {number}\n",
                    encoding="utf-8",
                )
            manifest = write_manifest(root)
            built, payload = run_json(
                str(BUILDER), "--manifest", str(manifest), "--changed-only", "--assemble"
            )
            self.assertEqual(built.returncode, 0, built.stderr)
            self.assertEqual(payload["manuscript"]["status"], "BUILT")
            manuscript = root / "final" / "manuscript.docx"
            self.assertTrue(manuscript.is_file())
            first_hash = file_hash(manuscript)

            skipped, skipped_payload = run_json(
                str(BUILDER), "--manifest", str(manifest), "--changed-only", "--assemble"
            )
            self.assertEqual(skipped.returncode, 0, skipped.stderr)
            self.assertEqual(skipped_payload["manuscript"]["status"], "SKIPPED")
            self.assertEqual(first_hash, file_hash(manuscript))

    def test_v2_prefers_revision_and_creates_only_final_docx(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            chapter = root / "chapters" / "chapter-01"
            chapter.mkdir(parents=True)
            draft = chapter / "draft.md"
            revised = chapter / "revised.md"
            draft.write_text("# บทที่ 1\n\nเนื้อหาร่างเดิม\n", encoding="utf-8")
            revised.write_text("# บทที่ 1\n\nเนื้อหาฉบับแก้ที่ผ่าน QC\n", encoding="utf-8")
            manifest = write_manifest_v2(root)

            rejected, rejected_payload = run_json(
                str(BUILDER), "--manifest", str(manifest)
            )
            self.assertEqual(rejected.returncode, 2)
            self.assertIn("--final-only", rejected_payload["error"])
            self.assertEqual(list(root.rglob("*.docx")), [])

            built, payload = run_json(
                str(BUILDER), "--manifest", str(manifest), "--final-only"
            )
            self.assertEqual(built.returncode, 0, built.stderr)
            self.assertEqual(payload["manuscript"]["status"], "BUILT")
            manuscript = root / "final" / "manuscript.docx"
            self.assertEqual(list(root.rglob("*.docx")), [manuscript])
            self.assertEqual(list(root.rglob("*.pdf")), [])
            text = "\n".join(p.text for p in Document(manuscript).paragraphs)
            self.assertIn("เนื้อหาฉบับแก้ที่ผ่าน QC", text)
            self.assertNotIn("เนื้อหาร่างเดิม", text)
            state = json.loads(
                (root / ".docx-build" / "state.json").read_text(encoding="utf-8")
            )
            self.assertTrue(state["chapters"]["chapter-01"]["source"].endswith("revised.md"))
            self.assertNotIn("output", state["chapters"]["chapter-01"])

            revised.unlink()
            rebuilt, rebuilt_payload = run_json(
                str(BUILDER),
                "--manifest",
                str(manifest),
                "--final-only",
                "--changed-only",
            )
            self.assertEqual(rebuilt.returncode, 0, rebuilt.stderr)
            self.assertEqual(rebuilt_payload["manuscript"]["status"], "BUILT")
            text = "\n".join(p.text for p in Document(manuscript).paragraphs)
            self.assertIn("เนื้อหาร่างเดิม", text)
            self.assertEqual(list(root.rglob("*.docx")), [manuscript])


class DocxPreflightTests(unittest.TestCase):
    def test_plain_chapter_passes_without_render_trigger(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            source = root / "chapters" / "chapter-01" / "draft.md"
            source.parent.mkdir(parents=True)
            source.write_text(
                "# บทที่ 1 บทนำ\n\n## 1.1 หลักการ\n\nเนื้อหาสำหรับตรวจโครงสร้าง\n\n- ประเด็นหนึ่ง\n",
                encoding="utf-8",
            )
            manifest = write_manifest(root)
            payload = yaml.safe_load(manifest.read_text(encoding="utf-8"))
            payload["chapters"] = [payload["chapters"][0]]
            manifest.write_text(yaml.safe_dump(payload, allow_unicode=True, sort_keys=False), encoding="utf-8")
            built, _ = run_json(str(BUILDER), "--manifest", str(manifest))
            self.assertEqual(built.returncode, 0, built.stderr)
            docx = root / "chapters" / "chapter-01" / "chapter-01.docx"
            audited, report = run_json(str(PREFLIGHT), str(docx))
            self.assertEqual(audited.returncode, 0, audited.stderr)
            self.assertEqual(report["status"], "PASS")
            self.assertEqual(report["render_reasons"], [])
            self.assertEqual(report["findings"], [])

    def test_tables_and_images_require_render_but_pass_structural_checks(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            source = root / "chapters" / "chapter-01" / "draft.md"
            source.parent.mkdir(parents=True)
            source.write_text(
                "# บทที่ 1 บทนำ\n\n## 1.1 สื่อประกอบ\n\n{{figure:framework}}\n\n{{table:comparison}}\n",
                encoding="utf-8",
            )
            assets = root / "assets"
            assets.mkdir()
            (assets / "framework.png").write_bytes(PNG_1X1)
            (assets / "comparison.csv").write_text("หัวข้อ,รายละเอียด\nก,ข้อมูล\n", encoding="utf-8")
            manifest = write_manifest(root, include_assets=True)
            built, _ = run_json(str(BUILDER), "--manifest", str(manifest))
            self.assertEqual(built.returncode, 0, built.stderr)
            docx = root / "chapters" / "chapter-01" / "chapter-01.docx"
            report_path = root / "audit.json"
            audited, report = run_json(str(PREFLIGHT), str(docx), "--output", str(report_path))
            self.assertEqual(audited.returncode, 0, audited.stderr)
            self.assertEqual(report["status"], "RENDER_REQUIRED")
            self.assertEqual(report["render_reasons"], ["IMAGES", "TABLES"])
            self.assertEqual(report["summary"]["tables"], 1)
            self.assertEqual(report["summary"]["images"], 1)
            self.assertEqual(report["findings"], [])
            self.assertEqual(report, json.loads(report_path.read_text(encoding="utf-8")))

    def test_invalid_package_returns_machine_readable_warning(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "broken.docx"
            with zipfile.ZipFile(path, "w") as archive:
                archive.writestr("word/document.xml", "<not-closed>")
            completed, report = run_json(str(PREFLIGHT), str(path))
            self.assertEqual(completed.returncode, 2)
            self.assertEqual(report["status"], "WARN")
            self.assertEqual(report["findings"][0]["code"], "PACKAGE_INVALID")


if __name__ == "__main__":
    unittest.main()

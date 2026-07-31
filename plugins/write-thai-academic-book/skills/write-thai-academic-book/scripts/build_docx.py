#!/usr/bin/env python3
"""Build deterministic DOCX files from a manifest and Markdown sources.

The builder keeps prose generation separate from document packaging. A single
manifest defines styles, chapter inputs, tables, images, outputs, and cache
locations. Manifest v2 is Markdown-first and permits only final manuscript
assembly; v1 remains readable for existing projects.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import sys
import tempfile
import zipfile
from copy import deepcopy
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

import yaml
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


BUILD_VERSION = "2"
FIXED_TIMESTAMP = datetime(2000, 1, 1, tzinfo=timezone.utc)
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"

DEFAULT_STYLE: dict[str, Any] = {
    "font": "TH Sarabun New",
    "body_size_pt": 16,
    "title_size_pt": 22,
    "heading_sizes_pt": [18, 16, 16],
    "caption_size_pt": 14,
    "line_spacing": 1.15,
    "first_line_indent_cm": 1.25,
    "page": {
        "size": "A4",
        "orientation": "portrait",
        "margin_top_cm": 3.0,
        "margin_bottom_cm": 2.5,
        "margin_left_cm": 3.0,
        "margin_right_cm": 2.5,
        "header_distance_cm": 1.25,
        "footer_distance_cm": 1.25,
    },
    "colors": {"text": "000000", "heading": "17365D", "muted": "666666"},
    "table": {"cell_margin_dxa": 120, "header_fill": "D9EAF7"},
}


class BuildError(RuntimeError):
    """Raised for a user-actionable manifest or build problem."""


def deep_merge(base: dict[str, Any], override: dict[str, Any]) -> dict[str, Any]:
    result = deepcopy(base)
    for key, value in override.items():
        if isinstance(value, dict) and isinstance(result.get(key), dict):
            result[key] = deep_merge(result[key], value)
        else:
            result[key] = value
    return result


def load_manifest(path: Path) -> dict[str, Any]:
    try:
        payload = yaml.safe_load(path.read_text(encoding="utf-8-sig"))
    except (OSError, yaml.YAMLError) as exc:
        raise BuildError(f"Cannot read manifest: {exc}") from exc
    if not isinstance(payload, dict):
        raise BuildError("Manifest root must be a mapping.")
    version = payload.get("version", 1)
    if version not in (1, 2):
        raise BuildError("Only manifest versions 1 and 2 are supported.")
    chapters = payload.get("chapters")
    if not isinstance(chapters, list) or not chapters:
        raise BuildError("Manifest must contain a non-empty chapters list.")
    if version == 2:
        for chapter in chapters:
            if not isinstance(chapter, dict):
                raise BuildError("Each chapter entry must be a mapping.")
            candidates = chapter.get("source_candidates")
            if (
                not isinstance(candidates, list)
                or not candidates
                or not all(isinstance(item, str) and item.strip() for item in candidates)
            ):
                raise BuildError(
                    "Manifest v2 chapters require a non-empty source_candidates list."
                )
            if chapter.get("output"):
                raise BuildError(
                    "Manifest v2 does not allow chapter output paths; use --final-only."
                )
    return payload


def resolve_path(base: Path, value: str | Path | None) -> Path | None:
    if value in (None, ""):
        return None
    path = Path(value)
    if not path.is_absolute():
        path = base / path
    return path.resolve()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def stable_hash(parts: Iterable[bytes]) -> str:
    digest = hashlib.sha256()
    for part in parts:
        digest.update(len(part).to_bytes(8, "big"))
        digest.update(part)
    return digest.hexdigest()


def json_bytes(value: Any) -> bytes:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")


def load_state(path: Path) -> dict[str, Any]:
    if not path.is_file():
        return {"version": 1, "chapters": {}, "manuscript": {}}
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return {"version": 1, "chapters": {}, "manuscript": {}}
    if not isinstance(payload, dict):
        return {"version": 1, "chapters": {}, "manuscript": {}}
    payload.setdefault("version", 1)
    payload.setdefault("chapters", {})
    payload.setdefault("manuscript", {})
    return payload


def write_json(path: Path, payload: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )


def color(value: str) -> RGBColor:
    value = value.lstrip("#")
    if not re.fullmatch(r"[0-9A-Fa-f]{6}", value):
        raise BuildError(f"Invalid RGB color: {value}")
    return RGBColor.from_string(value.upper())


def set_style_font(style: Any, font: str, size_pt: float, bold: bool = False, color_hex: str = "000000") -> None:
    style.font.name = font
    style.font.size = Pt(size_pt)
    style.font.bold = bold
    style.font.color.rgb = color(color_hex)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), font)


def get_or_add_style(doc: Document, name: str, style_type: WD_STYLE_TYPE) -> Any:
    try:
        return doc.styles[name]
    except KeyError:
        return doc.styles.add_style(name, style_type)


def configure_styles(doc: Document, style_config: dict[str, Any]) -> None:
    font = str(style_config["font"])
    body_size = float(style_config["body_size_pt"])
    heading_sizes = list(style_config["heading_sizes_pt"])
    while len(heading_sizes) < 3:
        heading_sizes.append(heading_sizes[-1] if heading_sizes else body_size)
    colors = style_config["colors"]
    line_spacing = float(style_config["line_spacing"])

    normal = doc.styles["Normal"]
    set_style_font(normal, font, body_size, color_hex=colors["text"])
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.first_line_indent = Cm(float(style_config["first_line_indent_cm"]))
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = line_spacing

    title = doc.styles["Title"]
    set_style_font(title, font, float(style_config["title_size_pt"]), bold=True, color_hex=colors["heading"])
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(14)
    title.paragraph_format.keep_with_next = True

    for index, size in enumerate(heading_sizes[:3], start=1):
        heading = doc.styles[f"Heading {index}"]
        set_style_font(heading, font, float(size), bold=True, color_hex=colors["heading"])
        heading.paragraph_format.space_before = Pt(14 if index == 1 else 10)
        heading.paragraph_format.space_after = Pt(6 if index == 1 else 4)
        heading.paragraph_format.keep_with_next = True
        heading.paragraph_format.keep_together = True

    body_no_indent = get_or_add_style(doc, "Body No Indent", WD_STYLE_TYPE.PARAGRAPH)
    set_style_font(body_no_indent, font, body_size, color_hex=colors["text"])
    body_no_indent.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    body_no_indent.paragraph_format.first_line_indent = Cm(0)
    body_no_indent.paragraph_format.space_after = Pt(0)
    body_no_indent.paragraph_format.line_spacing = line_spacing

    bibliography = get_or_add_style(doc, "Bibliography", WD_STYLE_TYPE.PARAGRAPH)
    set_style_font(bibliography, font, body_size, color_hex=colors["text"])
    bibliography.paragraph_format.left_indent = Cm(0.75)
    bibliography.paragraph_format.first_line_indent = Cm(-0.75)
    bibliography.paragraph_format.space_after = Pt(2)
    bibliography.paragraph_format.line_spacing = 1.0

    caption = get_or_add_style(doc, "Caption", WD_STYLE_TYPE.PARAGRAPH)
    set_style_font(caption, font, float(style_config["caption_size_pt"]), color_hex=colors["muted"])
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.keep_with_next = True

    code_style = get_or_add_style(doc, "Code Block", WD_STYLE_TYPE.PARAGRAPH)
    set_style_font(code_style, "Consolas", max(10, body_size - 2), color_hex=colors["text"])
    code_style.paragraph_format.left_indent = Cm(0.75)
    code_style.paragraph_format.space_before = Pt(3)
    code_style.paragraph_format.space_after = Pt(3)
    code_style.paragraph_format.line_spacing = 1.0

    for name in ("Header", "Footer", "List Bullet", "List Number"):
        try:
            target = doc.styles[name]
        except KeyError:
            continue
        set_style_font(target, font, body_size if name.startswith("List") else max(10, body_size - 2), color_hex=colors["muted"] if name in {"Header", "Footer"} else colors["text"])


def clear_template_body(doc: Document) -> None:
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def configure_page(doc: Document, style_config: dict[str, Any]) -> None:
    page = style_config["page"]
    for section in doc.sections:
        size = str(page.get("size", "A4")).upper()
        if size != "A4":
            raise BuildError(f"Unsupported page size: {size}; only A4 is currently supported.")
        orientation = str(page.get("orientation", "portrait")).lower()
        if orientation == "landscape":
            section.orientation = WD_ORIENT.LANDSCAPE
            section.page_width = Cm(29.7)
            section.page_height = Cm(21.0)
        else:
            section.orientation = WD_ORIENT.PORTRAIT
            section.page_width = Cm(21.0)
            section.page_height = Cm(29.7)
        section.top_margin = Cm(float(page["margin_top_cm"]))
        section.bottom_margin = Cm(float(page["margin_bottom_cm"]))
        section.left_margin = Cm(float(page["margin_left_cm"]))
        section.right_margin = Cm(float(page["margin_right_cm"]))
        section.header_distance = Cm(float(page["header_distance_cm"]))
        section.footer_distance = Cm(float(page["footer_distance_cm"]))


def add_page_field(paragraph: Any) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, text, end])


def configure_header_footer(doc: Document, project: dict[str, Any], style_config: dict[str, Any]) -> None:
    header_text = project.get("header_text") or project.get("title") or ""
    show_page = project.get("footer_page_numbers", True)
    for section in doc.sections:
        header = section.header.paragraphs[0]
        header.style = doc.styles["Header"]
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if header_text:
            header.add_run(str(header_text))
        footer = section.footer.paragraphs[0]
        footer.style = doc.styles["Footer"]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if show_page:
            add_page_field(footer)


def set_cell_margins(cell: Any, value: int) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge in ("top", "start", "bottom", "end"):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_shading(cell: Any, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_repeat_table_header(row: Any) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


def set_table_geometry(table: Any, widths_dxa: list[int], indent_dxa: int = 120) -> None:
    total = sum(widths_dxa)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    old_grid = table._tbl.tblGrid
    new_grid = OxmlElement("w:tblGrid")
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        new_grid.append(col)
    table._tbl.replace(old_grid, new_grid)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths_dxa[min(index, len(widths_dxa) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")


def usable_width_dxa(style_config: dict[str, Any]) -> int:
    page = style_config["page"]
    width_cm = 29.7 if str(page.get("orientation", "portrait")).lower() == "landscape" else 21.0
    usable_cm = width_cm - float(page["margin_left_cm"]) - float(page["margin_right_cm"])
    return max(1440, round(usable_cm / 2.54 * 1440))


def normalize_widths(column_count: int, width_values: list[float] | None, total_dxa: int) -> list[int]:
    if column_count <= 0:
        raise BuildError("Cannot create a table with zero columns.")
    if not width_values:
        base = total_dxa // column_count
        widths = [base] * column_count
    else:
        if len(width_values) != column_count or any(value <= 0 for value in width_values):
            raise BuildError("Table widths must contain one positive value per column.")
        scale = total_dxa / sum(width_values)
        widths = [round(value * scale) for value in width_values]
    widths[-1] += total_dxa - sum(widths)
    return widths


def add_hyperlink(paragraph: Any, text: str, url: str) -> None:
    rel_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    run_pr = OxmlElement("w:rPr")
    run_style = OxmlElement("w:rStyle")
    run_style.set(qn("w:val"), "Hyperlink")
    run_pr.append(run_style)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.extend([run_pr, text_node])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


INLINE_PATTERN = re.compile(
    r"(\*\*[^*]+\*\*|(?<!\*)\*[^*]+\*(?!\*)|`[^`]+`|\[[^\]]+\]\([^)]+\))"
)


def add_inline(paragraph: Any, text: str) -> None:
    position = 0
    for match in INLINE_PATTERN.finditer(text):
        if match.start() > position:
            paragraph.add_run(text[position : match.start()])
        token = match.group(0)
        if token.startswith("**"):
            paragraph.add_run(token[2:-2]).bold = True
        elif token.startswith("*"):
            paragraph.add_run(token[1:-1]).italic = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
        else:
            link = re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", token)
            if link:
                add_hyperlink(paragraph, link.group(1), link.group(2))
        position = match.end()
    if position < len(text):
        paragraph.add_run(text[position:])


def markdown_table(lines: list[str]) -> list[list[str]] | None:
    if len(lines) < 2:
        return None
    separator = [part.strip() for part in lines[1].strip().strip("|").split("|")]
    if not separator or not all(re.fullmatch(r":?-{3,}:?", cell) for cell in separator):
        return None
    rows = []
    for index, line in enumerate(lines):
        if index == 1:
            continue
        rows.append([part.strip() for part in line.strip().strip("|").split("|")])
    width = max(len(row) for row in rows)
    return [row + [""] * (width - len(row)) for row in rows]


def read_table(path: Path) -> list[list[str]]:
    suffix = path.suffix.lower()
    if suffix in {".csv", ".tsv"}:
        delimiter = "\t" if suffix == ".tsv" else ","
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            return [[str(cell) for cell in row] for row in csv.reader(handle, delimiter=delimiter)]
    if suffix == ".json":
        payload = json.loads(path.read_text(encoding="utf-8-sig"))
        if isinstance(payload, list) and payload and all(isinstance(item, dict) for item in payload):
            headers = list(payload[0])
            return [headers] + [[str(item.get(header, "")) for header in headers] for item in payload]
        if isinstance(payload, list) and all(isinstance(item, list) for item in payload):
            return [[str(cell) for cell in row] for row in payload]
        raise BuildError(f"JSON table must be a list of rows or objects: {path}")
    if suffix in {".md", ".markdown"}:
        lines = [line for line in path.read_text(encoding="utf-8-sig").splitlines() if line.strip()]
        parsed = markdown_table(lines)
        if parsed:
            return parsed
    raise BuildError(f"Unsupported or invalid table source: {path}")


def add_table(doc: Document, rows: list[list[str]], style_config: dict[str, Any], asset: dict[str, Any] | None = None) -> None:
    if not rows or not rows[0]:
        raise BuildError("Table data is empty.")
    column_count = max(len(row) for row in rows)
    normalized = [row + [""] * (column_count - len(row)) for row in rows]
    if asset and asset.get("caption"):
        caption = doc.add_paragraph(style="Caption")
        add_inline(caption, str(asset["caption"]))
    table = doc.add_table(rows=len(normalized), cols=column_count)
    table.style = "Table Grid"
    table_config = style_config["table"]
    for row_index, row in enumerate(normalized):
        for column_index, value in enumerate(row):
            cell = table.cell(row_index, column_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, int(table_config["cell_margin_dxa"]))
            paragraph = cell.paragraphs[0]
            paragraph.style = doc.styles["Body No Indent"]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_inline(paragraph, value)
            if row_index == 0:
                set_cell_shading(cell, str(table_config["header_fill"]))
                for run in paragraph.runs:
                    run.bold = True
    set_repeat_table_header(table.rows[0])
    widths = normalize_widths(
        column_count,
        [float(value) for value in asset.get("widths", [])] if asset and asset.get("widths") else None,
        usable_width_dxa(style_config),
    )
    set_table_geometry(table, widths, int(table_config["cell_margin_dxa"]))
    spacer = doc.add_paragraph(style="Body No Indent")
    spacer.paragraph_format.space_after = Pt(0)


def set_image_alt_text(shape: Any, alt_text: str, title: str = "") -> None:
    doc_pr = shape._inline.docPr
    doc_pr.set("descr", alt_text)
    if title:
        doc_pr.set("title", title)


def add_image(doc: Document, path: Path, alt_text: str, caption: str | None, width_cm: float | None) -> None:
    paragraph = doc.add_paragraph(style="Body No Indent")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    shape = run.add_picture(str(path), width=Cm(width_cm) if width_cm else None)
    set_image_alt_text(shape, alt_text or path.stem, title=path.stem)
    paragraph.paragraph_format.keep_with_next = bool(caption)
    if caption:
        caption_paragraph = doc.add_paragraph(style="Caption")
        add_inline(caption_paragraph, caption)


def assets_by_id(chapter: dict[str, Any], base: Path) -> tuple[dict[str, dict[str, Any]], dict[str, dict[str, Any]]]:
    assets = chapter.get("assets") or {}
    if not isinstance(assets, dict):
        raise BuildError(f"Chapter {chapter.get('id')} assets must be a mapping.")
    images: dict[str, dict[str, Any]] = {}
    tables: dict[str, dict[str, Any]] = {}
    for kind, target in (("images", images), ("tables", tables)):
        values = assets.get(kind) or []
        if not isinstance(values, list):
            raise BuildError(f"Chapter {chapter.get('id')} assets.{kind} must be a list.")
        for item in values:
            if not isinstance(item, dict) or not item.get("id") or not item.get("path"):
                raise BuildError(f"Each {kind} item requires id and path.")
            resolved = dict(item)
            resolved["path"] = str(resolve_path(base, item["path"]))
            target[str(item["id"])] = resolved
    return images, tables


def render_markdown(
    doc: Document,
    source: Path,
    chapter: dict[str, Any],
    manifest_base: Path,
    style_config: dict[str, Any],
) -> None:
    text = source.read_text(encoding="utf-8-sig")
    lines = text.splitlines()
    images, tables = assets_by_id(chapter, manifest_base)
    in_code = False
    code_lines: list[str] = []
    bibliography = False
    first_body_after_heading = False
    index = 0

    while index < len(lines):
        raw = lines[index]
        stripped = raw.strip()
        if stripped.startswith("```"):
            if in_code:
                paragraph = doc.add_paragraph(style="Code Block")
                paragraph.add_run("\n".join(code_lines))
                code_lines = []
                in_code = False
            else:
                in_code = True
            index += 1
            continue
        if in_code:
            code_lines.append(raw)
            index += 1
            continue
        if not stripped:
            index += 1
            continue

        if stripped.startswith("|"):
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            parsed = markdown_table(table_lines)
            if parsed:
                add_table(doc, parsed, style_config)
                continue
            for line in table_lines:
                paragraph = doc.add_paragraph(style="Normal")
                add_inline(paragraph, line)
            continue

        marker = re.fullmatch(r"\{\{(figure|table):([A-Za-z0-9_.-]+)\}\}", stripped)
        if marker:
            kind, asset_id = marker.groups()
            asset = images.get(asset_id) if kind == "figure" else tables.get(asset_id)
            if asset is None:
                raise BuildError(f"Unknown {kind} asset '{asset_id}' in {source}")
            asset_path = Path(asset["path"])
            if not asset_path.is_file():
                raise BuildError(f"Missing asset: {asset_path}")
            if kind == "figure":
                add_image(
                    doc,
                    asset_path,
                    str(asset.get("alt", asset_path.stem)),
                    str(asset["caption"]) if asset.get("caption") else None,
                    float(asset["width_cm"]) if asset.get("width_cm") else None,
                )
            else:
                add_table(doc, read_table(asset_path), style_config, asset)
            index += 1
            continue

        direct_image = re.fullmatch(r"!\[([^\]]*)\]\(([^)]+)\)", stripped)
        if direct_image:
            image_path = resolve_path(source.parent, direct_image.group(2))
            if image_path is None or not image_path.is_file():
                raise BuildError(f"Missing image referenced by Markdown: {direct_image.group(2)}")
            add_image(doc, image_path, direct_image.group(1) or image_path.stem, None, None)
            index += 1
            continue

        if stripped == "{{pagebreak}}":
            paragraph = doc.add_paragraph(style="Body No Indent")
            paragraph.add_run().add_break(WD_BREAK.PAGE)
            index += 1
            continue

        heading = re.match(r"^(#{1,4})\s+(.+)$", stripped)
        if heading:
            level = len(heading.group(1))
            title = heading.group(2).strip()
            if level == 1:
                paragraph = doc.add_paragraph(style="Title")
            else:
                paragraph = doc.add_paragraph(style=f"Heading {min(level - 1, 3)}")
            add_inline(paragraph, title)
            bibliography = bool(re.search(r"บรรณานุกรม|เอกสารอ้างอิง|references", title, re.IGNORECASE))
            first_body_after_heading = True
            index += 1
            continue

        unordered = re.match(r"^[-*+]\s+(.+)$", stripped)
        ordered = re.match(r"^\d+[.)]\s+(.+)$", stripped)
        if unordered or ordered:
            paragraph = doc.add_paragraph(style="List Bullet" if unordered else "List Number")
            add_inline(paragraph, (unordered or ordered).group(1))
            index += 1
            continue

        if stripped.startswith("> "):
            paragraph = doc.add_paragraph(style="Quote")
            add_inline(paragraph, stripped[2:])
            index += 1
            continue

        style_name = "Bibliography" if bibliography else ("Body No Indent" if first_body_after_heading else "Normal")
        paragraph = doc.add_paragraph(style=style_name)
        add_inline(paragraph, stripped)
        first_body_after_heading = False
        index += 1

    if in_code:
        paragraph = doc.add_paragraph(style="Code Block")
        paragraph.add_run("\n".join(code_lines))


def apply_core_properties(doc: Document, metadata: dict[str, Any]) -> None:
    core = doc.core_properties
    core.title = str(metadata.get("title", ""))
    core.subject = str(metadata.get("subject", ""))
    core.author = str(metadata.get("author", ""))
    keywords = metadata.get("keywords", "")
    core.keywords = ", ".join(str(item) for item in keywords) if isinstance(keywords, list) else str(keywords)
    core.comments = str(metadata.get("description", "generated by build_docx.py"))
    core.created = FIXED_TIMESTAMP
    core.modified = FIXED_TIMESTAMP
    core.revision = 1


def normalized_docx(source: Path, destination: Path) -> None:
    destination.parent.mkdir(parents=True, exist_ok=True)
    rsid_pattern = re.compile(rb"\s+w:rsid[A-Za-z]+\s*=\s*\"[^\"]*\"")
    with zipfile.ZipFile(source, "r") as archive, zipfile.ZipFile(
        destination, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9
    ) as output:
        for name in sorted(archive.namelist()):
            data = archive.read(name)
            if name.endswith(".xml") or name.endswith(".rels"):
                data = rsid_pattern.sub(b"", data)
            info = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.create_system = 0
            info.external_attr = 0
            output.writestr(info, data)


def create_document(template: Path | None, project: dict[str, Any], style_config: dict[str, Any]) -> Document:
    if template:
        if not template.is_file():
            raise BuildError(f"Style template not found: {template}")
        doc = Document(str(template))
        clear_template_body(doc)
    else:
        doc = Document()
    configure_page(doc, style_config)
    configure_styles(doc, style_config)
    configure_header_footer(doc, project, style_config)
    return doc


def build_document(
    output: Path,
    chapters: list[tuple[dict[str, Any], Path]],
    manifest_base: Path,
    project: dict[str, Any],
    style_config: dict[str, Any],
    template: Path | None,
    manuscript: bool = False,
) -> None:
    doc = create_document(template, project, style_config)
    for index, (chapter, source) in enumerate(chapters):
        if index and manuscript:
            doc.add_page_break()
        render_markdown(doc, source, chapter, manifest_base, style_config)
    metadata = deep_merge(project, chapters[0][0].get("metadata") or {}) if len(chapters) == 1 else project
    apply_core_properties(doc, metadata)
    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="build-docx-") as temp_dir:
        intermediate = Path(temp_dir) / "document.docx"
        doc.save(intermediate)
        normalized_docx(intermediate, output)


def asset_paths(chapter: dict[str, Any], manifest_base: Path, source: Path) -> list[Path]:
    images, tables = assets_by_id(chapter, manifest_base)
    paths = [Path(item["path"]) for item in [*images.values(), *tables.values()]]
    for match in re.finditer(r"!\[[^\]]*\]\(([^)]+)\)", source.read_text(encoding="utf-8-sig")):
        direct = resolve_path(source.parent, match.group(1))
        if direct:
            paths.append(direct)
    return sorted(set(paths), key=lambda path: str(path).lower())


def chapter_digest(
    chapter: dict[str, Any],
    source: Path,
    manifest_base: Path,
    project: dict[str, Any],
    style_config: dict[str, Any],
    template: Path | None,
) -> tuple[str, list[Path]]:
    dependencies = [source, *asset_paths(chapter, manifest_base, source)]
    if template:
        dependencies.append(template)
    for dependency in dependencies:
        if not dependency.is_file():
            raise BuildError(f"Missing build dependency: {dependency}")
    config = {
        "builder": BUILD_VERSION,
        "project": project,
        "style": style_config,
        "chapter": chapter,
    }
    parts = [json_bytes(config)]
    for dependency in dependencies:
        parts.extend([str(dependency).encode("utf-8"), dependency.read_bytes()])
    return stable_hash(parts), dependencies


def output_for_chapter(chapter: dict[str, Any], manifest_base: Path, build_dir: Path) -> Path:
    chapter_id = str(chapter.get("id") or "").strip()
    if not chapter_id:
        raise BuildError("Every chapter requires a non-empty id.")
    return resolve_path(manifest_base, chapter.get("output")) or (build_dir / f"{chapter_id}.docx")


def source_for_chapter(
    chapter: dict[str, Any], manifest_base: Path, manifest_version: int
) -> Path:
    chapter_id = str(chapter.get("id") or "").strip()
    if manifest_version == 1:
        source = resolve_path(manifest_base, chapter.get("source"))
        if source is None or not source.is_file():
            raise BuildError(f"Chapter {chapter_id} source not found: {source}")
        return source

    candidates = chapter.get("source_candidates") or []
    resolved = [resolve_path(manifest_base, candidate) for candidate in candidates]
    for source in resolved:
        if source is not None and source.is_file():
            return source
    choices = ", ".join(str(path) for path in resolved)
    raise BuildError(
        f"Chapter {chapter_id} has no existing Markdown source candidate: {choices}"
    )


def build_from_manifest(
    manifest_path: Path,
    changed_only: bool,
    assemble: bool,
    final_only: bool = False,
) -> dict[str, Any]:
    manifest = load_manifest(manifest_path)
    manifest_version = int(manifest.get("version", 1))
    if manifest_version == 2 and not final_only:
        raise BuildError("Manifest v2 must be built with --final-only.")
    if final_only:
        assemble = True
    base = manifest_path.parent.resolve()
    project = manifest.get("project") or {}
    if not isinstance(project, dict):
        raise BuildError("project must be a mapping.")
    style_config = deep_merge(DEFAULT_STYLE, manifest.get("style") or {})
    template = resolve_path(base, style_config.get("template"))
    build_config = manifest.get("build") or {}
    build_dir = resolve_path(base, build_config.get("directory", ".docx-build"))
    assert build_dir is not None
    state_path = resolve_path(base, build_config.get("state", str(build_dir / "state.json")))
    assert state_path is not None
    state = load_state(state_path)
    prior_chapters = state.get("chapters", {})
    next_chapters: dict[str, Any] = {}
    resolved_chapters: list[tuple[dict[str, Any], Path, Path | None, str]] = []
    rebuilt: list[str] = []
    skipped: list[str] = []

    seen_ids: set[str] = set()
    for chapter in manifest["chapters"]:
        if not isinstance(chapter, dict):
            raise BuildError("Each chapter entry must be a mapping.")
        chapter_id = str(chapter.get("id") or "").strip()
        if not chapter_id or chapter_id in seen_ids:
            raise BuildError(f"Chapter ids must be unique and non-empty: {chapter_id!r}")
        seen_ids.add(chapter_id)
        source = source_for_chapter(chapter, base, manifest_version)
        output = None if final_only else output_for_chapter(chapter, base, build_dir)
        digest, dependencies = chapter_digest(chapter, source, base, project, style_config, template)
        previous = prior_chapters.get(chapter_id) if isinstance(prior_chapters, dict) else None
        if final_only:
            unchanged = bool(
                changed_only
                and isinstance(previous, dict)
                and previous.get("digest") == digest
            )
            (skipped if unchanged else rebuilt).append(chapter_id)
        else:
            assert output is not None
            unchanged = bool(
                changed_only
                and isinstance(previous, dict)
                and previous.get("digest") == digest
                and output.is_file()
            )
            if unchanged:
                skipped.append(chapter_id)
            else:
                build_document(output, [(chapter, source)], base, project, style_config, template)
                rebuilt.append(chapter_id)
        chapter_state = {
            "digest": digest,
            "source": str(source),
            "dependencies": [str(path) for path in dependencies],
        }
        if output is not None:
            chapter_state.update(
                {"output": str(output), "output_sha256": sha256_file(output)}
            )
        next_chapters[chapter_id] = chapter_state
        resolved_chapters.append((chapter, source, output, digest))

    manuscript_result: dict[str, Any] = {"status": "NOT_REQUESTED"}
    if assemble:
        manuscript_output = resolve_path(base, build_config.get("manuscript_output"))
        if manuscript_output is None:
            raise BuildError("build.manuscript_output is required with --assemble.")
        if manifest_version == 2:
            required_output = (base / "final" / "manuscript.docx").resolve()
            if manuscript_output != required_output:
                raise BuildError(
                    "Manifest v2 build.manuscript_output must be final/manuscript.docx."
                )
        manuscript_digest = stable_hash(
            [
                json_bytes({"builder": BUILD_VERSION, "project": project, "style": style_config}),
                *[digest.encode("ascii") for _, _, _, digest in resolved_chapters],
            ]
        )
        previous = state.get("manuscript") if isinstance(state.get("manuscript"), dict) else {}
        unchanged = bool(
            changed_only
            and previous.get("digest") == manuscript_digest
            and manuscript_output.is_file()
        )
        if not unchanged:
            build_document(
                manuscript_output,
                [(chapter, source) for chapter, source, _, _ in resolved_chapters],
                base,
                project,
                style_config,
                template,
                manuscript=True,
            )
        manuscript_result = {
            "status": "SKIPPED" if unchanged else "BUILT",
            "digest": manuscript_digest,
            "output": str(manuscript_output),
            "output_sha256": sha256_file(manuscript_output),
        }

    next_state = {
        "version": 2,
        "builder_version": BUILD_VERSION,
        "manifest": str(manifest_path.resolve()),
        "chapters": next_chapters,
        "manuscript": manuscript_result if assemble else state.get("manuscript", {}),
    }
    write_json(state_path, next_state)
    return {
        "status": "BUILT" if rebuilt or manuscript_result.get("status") == "BUILT" else "SKIPPED",
        "manifest": str(manifest_path.resolve()),
        "state": str(state_path),
        "rebuilt_chapters": rebuilt,
        "skipped_chapters": skipped,
        "manuscript": manuscript_result,
    }


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Build legacy chapter DOCX files or a v2 final-only manuscript."
    )
    parser.add_argument("--manifest", required=True, type=Path, help="Path to document.yml")
    parser.add_argument(
        "--changed-only",
        action="store_true",
        help="Reuse chapter outputs whose source, assets, template, and config digest is unchanged.",
    )
    parser.add_argument(
        "--assemble",
        action="store_true",
        help="Legacy v1: also assemble all chapters into build.manuscript_output.",
    )
    parser.add_argument(
        "--final-only",
        action="store_true",
        help="Assemble final/manuscript.docx directly from Markdown without chapter DOCX files.",
    )
    return parser.parse_args()


def main() -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    args = parse_args()
    try:
        result = build_from_manifest(
            args.manifest.resolve(),
            args.changed_only,
            args.assemble,
            args.final_only,
        )
    except (BuildError, OSError, ValueError, json.JSONDecodeError) as exc:
        print(json.dumps({"status": "ERROR", "error": str(exc)}, ensure_ascii=False, indent=2))
        return 2
    print(json.dumps(result, ensure_ascii=False, indent=2, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
